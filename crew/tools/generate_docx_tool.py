from docx import Document
from bs4 import BeautifulSoup
import os
import markdown
from crewai.tools import BaseTool

class GenerateDocxTool(BaseTool):
    name: str = "GenerateDocx"
    description: str = "Gera um documento Word a partir de um Markdown."
    word_file: str

    def _run(self, markdown_text:str) -> str:
        try:
            # Criar diretório de relatórios se não existir
            relatorios_dir = os.path.join(os.getcwd(), 'relatorios')
            # relatorios_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "relatorios")
            os.makedirs(relatorios_dir, exist_ok=True)

            # Converting Markdown to HTML
            html_content = markdown.markdown(markdown_text)

            # Creating a new Word Document
            doc = Document()

            # Converting HTML to text and add it to the Word Document
            soup = BeautifulSoup(html_content, 'html.parser')
            
            # Adding content to the Word Document
            for element in soup:
                if element.name == 'h1':
                    doc.add_heading(element.text, level=1)
                elif element.name == 'h2':
                    doc.add_heading(element.text, level=2)
                elif element.name == 'h3':
                    doc.add_heading(element.text, level=3)
                elif element.name == 'p':
                    paragraph = doc.add_paragraph()
                    for child in element.children:
                        if child.name == 'strong':
                            paragraph.add_run(child.text).bold = True
                        elif child.name == 'em':
                            paragraph.add_run(child.text).italic = True
                        else:
                            paragraph.add_run(child)
                elif element.name == 'ul':
                    for li in element.find_all('li'):
                        doc.add_paragraph(li.text, style='List Bullet')
                elif element.name == 'ol':
                    for li in element.find_all('li'):
                        doc.add_paragraph(li.text, style='List Number')
            
            docx_name = os.path.join(relatorios_dir, f'{self.word_file}.docx')
            doc.save(docx_name)
            print(f'Documento {docx_name} gerado com sucesso!')
            print(docx_name)
            return docx_name
        except Exception as e:
            raise e